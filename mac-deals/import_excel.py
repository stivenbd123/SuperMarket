# Importa la hoja "Precios por modelo" de tu Excel a la tabla market_refs.
# Se corre a mano cada vez que actualices el Excel:
#   docker compose run --rm app python import_excel.py
from openpyxl import load_workbook

import db

EXCEL_FILE = "Macs_Marketplace_Colombia_oct2026.xlsx"
SHEET = "Precios por modelo"


def to_int(value):
    # Celda vacía -> None, número -> entero
    if value is None or value == "":
        return None
    return int(round(float(value)))


conn = db.connect()
workbook = load_workbook(EXCEL_FILE, data_only=True)  # data_only: el valor, no la fórmula
sheet = workbook[SHEET]

count = 0
# min_row=2 para saltar los encabezados
for row in sheet.iter_rows(min_row=2, values_only=True):
    group = row[0]       # A
    if not group:
        continue         # fila vacía

    conn.execute(
        """INSERT INTO market_refs
             (group_key, listings_count, min_cop, p25_cop, median_cop, max_cop,
              my_sale_cop, imported_at)
           VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
           ON CONFLICT (group_key) DO UPDATE SET
             listings_count = excluded.listings_count,
             min_cop        = excluded.min_cop,
             p25_cop        = excluded.p25_cop,
             median_cop     = excluded.median_cop,
             max_cop        = excluded.max_cop,
             my_sale_cop    = excluded.my_sale_cop,
             imported_at    = excluded.imported_at""",
        (str(group).strip(), to_int(row[1]), to_int(row[2]), to_int(row[3]),
         to_int(row[4]), to_int(row[5]), to_int(row[6]), db.now()),
    )
    count += 1

conn.commit()
print(count, "grupos importados")
